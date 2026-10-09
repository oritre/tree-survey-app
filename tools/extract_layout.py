# -*- coding: utf-8 -*-
"""
מחלץ מהתבנית (template/survey.xltm) את מה שצריך כדי לצייר את הדוח להדפסה ב-PDF
בתוך האפליקציה: רוחבי עמודות, גובהי שורות, שורות מוסתרות וצהובות, תאים ממוזגים,
טקסטים קבועים וסגנונות (גופן, יישור, מסגרות). הפלט: template/layout.json,
ובנוסף letterhead.png (נייר המכתבים מהכותרת), footer.png ו-stamp.png.

מריצים מחדש רק אם התבנית משתנה:  python3 tools/extract_layout.py
"""
import io
import json
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import get_column_letter
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / 'template' / 'survey.xltm'
OUT = ROOT / 'template'
LAST_ROW = 278
AREAS = [(1, 9), (11, 14)]  # A:I, K:N  (אזור ההדפסה בתבנית)
CHECK_COLUMNS = 14  # שורה צהובה: לפחות חצי מהתאים A-N צבועים צהוב (כמו בתוכנה במחשב)
DEFAULT_COL_WIDTH = 9.05  # רוחב ברירת מחדל בתווים, נמדד מ-PDF אמיתי של אקסל


def is_yellow(c):
    f = c.fill
    if not f or f.fill_type != 'solid':
        return False
    col = f.fgColor
    if col.type == 'rgb':
        return str(col.rgb).upper().endswith('FFFF00')
    return col.type == 'indexed' and col.indexed == 13


THEME_COLORS = {0: 'FFFFFF', 1: '000000', 2: 'E7E6E6', 3: '44546A'}


def color_of(c):
    if c is None:
        return None
    if c.type == 'rgb' and isinstance(c.rgb, str) and len(c.rgb) == 8:
        return c.rgb[2:]
    if c.type == 'theme' and c.theme in THEME_COLORS:
        return THEME_COLORS[c.theme]
    if c.type == 'indexed' and c.indexed == 13:
        return 'FFFF00'
    return None


def font_of(f):
    d = {}
    if f.sz:
        d['sz'] = float(f.sz)
    if f.b:
        d['b'] = 1
    if f.i:
        d['i'] = 1
    if f.u:
        d['u'] = 1
    col = color_of(f.color)
    if col and col != '000000':
        d['c'] = col
    return d


def main():
    wb = openpyxl.load_workbook(TEMPLATE, rich_text=True)
    ws = wb.worksheets[0]

    anchor = {}
    merges = []
    for m in ws.merged_cells.ranges:
        merges.append([m.min_row, m.min_col, m.max_row, m.max_col])
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                anchor[(r, c)] = (m.min_row, m.min_col)

    yellow = [r for r in range(1, LAST_ROW + 1)
              if sum(is_yellow(ws.cell(*anchor.get((r, c), (r, c)))) for c in range(1, CHECK_COLUMNS + 1)) * 2 >= CHECK_COLUMNS]

    cols = {c: DEFAULT_COL_WIDTH for c in range(1, 15)}
    for dim in ws.column_dimensions.values():  # openpyxl שומר טווח עמודות (min..max) תחת האות הראשונה
        if dim.customWidth and dim.width and dim.min:
            for c in range(dim.min, (dim.max or dim.min) + 1):
                if c in cols:
                    cols[c] = dim.width

    rows = {}
    for r in range(1, LAST_ROW + 1):
        dim = ws.row_dimensions.get(r)
        h = dim.height if dim is not None and dim.height else ws.sheet_format.defaultRowHeight
        rows[r] = [h, 1 if dim is not None and dim.hidden else 0]

    styles, style_ix = [], {}

    def style_id(cell):
        al = cell.alignment
        bd = cell.border
        s = {'f': font_of(cell.font)}
        if al.horizontal and al.horizontal != 'general':
            s['h'] = al.horizontal
        if al.vertical:
            s['v'] = al.vertical
        if al.wrap_text:
            s['w'] = 1
        if al.indent:
            s['ind'] = al.indent
        b = {}
        for side in ('left', 'right', 'top', 'bottom'):
            e = getattr(bd, side)
            if e is not None and e.style:
                b[side[0]] = [e.style, color_of(e.color) or '000000']
        if b:
            s['b'] = b
        fill = cell.fill
        if fill and fill.fill_type == 'solid':
            fc = color_of(fill.fgColor)
            if fc and fc != 'FFFFFF':
                s['bg'] = fc
        if cell.number_format and cell.number_format != 'General':
            s['nf'] = cell.number_format
        key = json.dumps(s, sort_keys=True)
        if key not in style_ix:
            style_ix[key] = len(styles)
            styles.append(s)
        return style_ix[key]

    cells = {}
    for c0, c1 in AREAS:
        for r in range(1, LAST_ROW + 1):
            for c in range(c0, c1 + 1):
                cell = ws.cell(r, c)
                v = cell.value
                entry = {}
                sid = style_id(cell)
                if sid:
                    entry['s'] = sid
                if isinstance(v, CellRichText):
                    runs = []
                    for part in v:
                        if isinstance(part, str):
                            runs.append([part, None])
                        else:
                            runs.append([part.text, font_of(part.font)])
                    entry['rt'] = runs
                elif isinstance(v, str) and v.startswith('='):
                    entry['fx'] = v
                elif v is not None and str(v) != '':
                    entry['t'] = v if isinstance(v, (int, float)) else str(v)
                if len(entry) > (1 if 's' in entry else 0) or 's' in entry:
                    cells[f'{get_column_letter(c)}{r}'] = entry

    # תמונות: נייר המכתבים בכותרת, פס תחתון, וחותמת שמעוגנת בגיליון
    z = zipfile.ZipFile(TEMPLATE)
    images = {}
    vml = z.read('xl/drawings/vmlDrawing1.vml').decode('utf-8', 'ignore')
    vrels = z.read('xl/drawings/_rels/vmlDrawing1.vml.rels').decode()
    for sid, style, rid in re.findall(r'<v:shape id="(\w+)".*?style=\'([^\']*)\'.*?o:relid="(\w+)"', vml, re.S):
        target = re.search(r'Id="%s"[^>]*Target="\.\./media/([^"]+)"' % rid, vrels).group(1)
        w = float(re.search(r'width:([\d.]+)pt', style).group(1))
        h = float(re.search(r'height:([\d.]+)pt', style).group(1))
        img = Image.open(io.BytesIO(z.read('xl/media/' + target))).convert('RGBA')
        name = {'LH': 'letterhead', 'CF': 'footer'}.get(sid, sid)
        if name == 'letterhead':
            # רק החלק העליון של נייר המכתבים מצויר (הלוגו), השאר שקוף
            bbox = img.getbbox()
            img = img.crop((0, 0, img.width, bbox[3] + 4))
            crop_h = h * img.height / Image.open(io.BytesIO(z.read('xl/media/' + target))).height
            img.thumbnail((1240, 10000))
            img.save(OUT / 'letterhead.png', optimize=True)
            images['letterhead'] = {'w': w, 'h': round(crop_h, 2)}
        else:
            img.save(OUT / (name + '.png'), optimize=True)
            images[name] = {'w': w, 'h': h}

    drawing = z.read('xl/drawings/drawing1.xml').decode()
    drels = z.read('xl/drawings/_rels/drawing1.xml.rels').decode()
    anchors = []
    for a in re.findall(r'<xdr:twoCellAnchor.*?</xdr:twoCellAnchor>', drawing, re.S):
        fr = re.search(r'<xdr:from><xdr:col>(\d+)</xdr:col><xdr:colOff>(\d+)</xdr:colOff><xdr:row>(\d+)</xdr:row><xdr:rowOff>(\d+)</xdr:rowOff>', a)
        ext = re.search(r'<a:ext cx="(\d+)" cy="(\d+)"', a)
        rid = re.search(r'r:embed="(\w+)"', a)
        if not (fr and ext and rid):
            continue
        target = re.search(r'Id="%s"[^>]*Target="\.\./media/([^"]+)"' % rid.group(1), drels).group(1)
        Image.open(io.BytesIO(z.read('xl/media/' + target))).save(OUT / 'stamp.png', optimize=True)
        anchors.append({'img': 'stamp', 'col': int(fr.group(1)) + 1, 'colOff': int(fr.group(2)) / 12700,
                        'row': int(fr.group(3)) + 1, 'rowOff': int(fr.group(4)) / 12700,
                        'w': int(ext.group(1)) / 12700, 'h': int(ext.group(2)) / 12700})

    sh = z.read('xl/worksheets/sheet1.xml').decode()
    pm = re.search(r'<pageMargins ([^/]*)/>', sh).group(1)
    margins = {k: float(v) * 72 for k, v in re.findall(r'(\w+)="([\d.]+)"', pm)}

    layout = {'cols': cols, 'rows': rows, 'yellow': yellow, 'merges': merges, 'styles': styles,
              'cells': cells, 'images': images, 'anchors': anchors, 'margins': margins, 'areas': AREAS,
              'lastRow': LAST_ROW}
    (OUT / 'layout.json').write_text(json.dumps(layout, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print('cells', len(cells), 'styles', len(styles), 'yellow', yellow, 'images', images, 'anchors', anchors)


if __name__ == '__main__':
    sys.exit(main())
